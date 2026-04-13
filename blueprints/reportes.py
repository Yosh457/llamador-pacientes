# blueprints/reportes.py
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from flask import Blueprint, jsonify, request, render_template, abort, send_file
from flask_login import login_required, current_user

from sqlalchemy import func
from datetime import datetime

from models import db, Llamado, LlamadoEvento, Box, Establecimiento
from utils import obtener_hora_chile

reportes_bp = Blueprint('reportes', __name__, template_folder='../templates', url_prefix='/reportes')

# --- PROTECCIÓN GLOBAL DEL BLUEPRINT ---
@reportes_bp.before_request
@login_required
def before_request():
    """Solo el Admin y el Supervisor pueden ver los reportes."""
    if current_user.rol.nombre not in ['Admin', 'Supervisor']:
        abort(403)

# --- FASE 4B: CASCARÓN HTML ---
@reportes_bp.route('/dashboard')
def dashboard():
    """
    Vista HTML del Dashboard Interactivo.
    """
    establecimientos = []
    # Si es Admin, le mandamos todos los establecimientos para el filtro
    if current_user.rol.nombre == 'Admin':
        establecimientos = Establecimiento.query.filter_by(activo=True).order_by(Establecimiento.nombre).all()
        
    return render_template(
        'reportes/dashboard.html', 
        establecimientos=establecimientos
    )

# --- FASE 4A: API DE MÉTRICAS ---
@reportes_bp.route('/api/metricas')
def api_metricas():
    """
    Endpoint JSON para consumo vía AJAX (Chart.js) o Postman.
    Calcula tiempos de espera, atención y distribución de estados.
    """
    # 1. Recibir Filtros (Query Params)
    fecha_desde_str = request.args.get('desde')
    fecha_hasta_str = request.args.get('hasta')
    est_filtro = request.args.get('establecimiento_id')

    query = Llamado.query

    # 2. Permisos y Filtros de Establecimiento
    if current_user.rol.nombre == 'Supervisor':
        # Supervisor está amarrado a su propio establecimiento
        query = query.filter_by(establecimiento_id=current_user.establecimiento_id)
    elif est_filtro and est_filtro.isdigit():
        query = query.filter_by(establecimiento_id=int(est_filtro))

    # 3. Filtros de Fecha (Por defecto: mes actual completo)
    ahora = obtener_hora_chile()
    
    if fecha_desde_str:
        fecha_desde = datetime.strptime(fecha_desde_str, '%Y-%m-%d').date()
    else:
        fecha_desde = ahora.replace(day=1).date() # Primer día del mes

    if fecha_hasta_str:
        fecha_hasta = datetime.strptime(fecha_hasta_str, '%Y-%m-%d').date()
    else:
        fecha_hasta = ahora.date() # Hoy

    # Aplicar filtro de fechas a la consulta
    query = query.filter(func.date(Llamado.fecha_creacion) >= fecha_desde)
    query = query.filter(func.date(Llamado.fecha_creacion) <= fecha_hasta)

    llamados = query.all()

    # --- INICIO DE CÁLCULOS MATEMÁTICOS OPTIMIZADOS ---
    total_llamados = len(llamados)
    
    # MEJORA: Estados dinámicos (Evita el hardcodeo y lo hace más flexible)
    estados = {}
    boxes_stats = {}

    tiempo_espera_total = 0
    llamados_con_espera = 0

    tiempo_atencion_total = 0
    llamados_con_atencion = 0

    total_rellamados = 0

    # MEJORA: Prevención de problema N+1 Queries. 
    # Traemos todos los eventos asociados a los llamados de este periodo en UNA sola consulta.
    eventos_por_llamado = {}
    if llamados:
        llamados_ids = [ll.id for ll in llamados]
        eventos_bulk = LlamadoEvento.query.filter(
            LlamadoEvento.llamado_id.in_(llamados_ids)
        ).order_by(LlamadoEvento.fecha_evento.asc()).all()

        for ev in eventos_bulk:
            eventos_por_llamado.setdefault(ev.llamado_id, []).append(ev)

    for ll in llamados:
        # Contadores de Estado dinámico
        estados[ll.estado] = estados.get(ll.estado, 0) + 1
        
        # MEJORA: Validación de Box (Null-safe) por si se eliminó el box o el dato está corrupto
        box_nombre = ll.box.nombre if ll.box else "Sin Box"
        boxes_stats[box_nombre] = boxes_stats.get(box_nombre, 0) + 1

        # Obtener eventos desde el diccionario en memoria O(1) en lugar de la BD O(N)
        eventos = eventos_por_llamado.get(ll.id, [])

        t_creacion = ll.fecha_creacion
        t_primer_llamado = None
        t_cierre = ll.fecha_cierre

        for ev in eventos:
            if ev.tipo_evento == 'PRIMER_LLAMADO':
                t_primer_llamado = ev.fecha_evento
            elif ev.tipo_evento in ['SEGUNDO_LLAMADO', 'TERCER_LLAMADO']:
                total_rellamados += 1

        # Calcular Tiempos (Solo si los eventos ocurrieron)
        if t_primer_llamado:
            # Espera = Creación (Ingreso de datos) -> Primer Llamado
            espera = (t_primer_llamado - t_creacion).total_seconds() / 60.0 # Convertir a minutos
            tiempo_espera_total += espera
            llamados_con_espera += 1

            # Atención = Primer Llamado -> Cierre Exitoso
            if t_cierre and ll.estado == 'FINALIZADO':
                atencion = (t_cierre - t_primer_llamado).total_seconds() / 60.0
                tiempo_atencion_total += atencion
                llamados_con_atencion += 1

    # --- PROMEDIOS FINALES ---
    promedio_espera = round(tiempo_espera_total / llamados_con_espera, 1) if llamados_con_espera > 0 else 0
    promedio_atencion = round(tiempo_atencion_total / llamados_con_atencion, 1) if llamados_con_atencion > 0 else 0
    promedio_rellamados = round(total_rellamados / total_llamados, 2) if total_llamados > 0 else 0

    # 4. Retornar la estructura JSON limpia
    return jsonify({
        "status": "success",
        "periodo": {
            "desde": fecha_desde.isoformat(),
            "hasta": fecha_hasta.isoformat()
        },
        "totales": {
            "general": total_llamados,
            "por_estado": estados,
            "por_box": boxes_stats
        },
        "tiempos_promedio_minutos": {
            "espera_creacion_a_primer_llamado": promedio_espera,
            "atencion_primer_llamado_a_cierre": promedio_atencion
        },
        "eficiencia": {
            "total_rellamados": total_rellamados,
            "promedio_rellamados_por_paciente": promedio_rellamados
        }
    })
    
# --- FASE 4C: EXPORTACIÓN A EXCEL ---
@reportes_bp.route('/export/excel')
def export_excel():
    """Genera y descarga un reporte detallado en formato Excel (.xlsx)."""
    
    # 1. Replicar la misma lógica de filtros de la API
    fecha_desde_str = request.args.get('desde')
    fecha_hasta_str = request.args.get('hasta')
    est_filtro = request.args.get('establecimiento_id')

    query = Llamado.query

    if current_user.rol.nombre == 'Supervisor':
        query = query.filter_by(establecimiento_id=current_user.establecimiento_id)
    elif est_filtro and est_filtro.isdigit():
        query = query.filter_by(establecimiento_id=int(est_filtro))

    ahora = obtener_hora_chile()
    
    if fecha_desde_str:
        fecha_desde = datetime.strptime(fecha_desde_str, '%Y-%m-%d').date()
    else:
        fecha_desde = ahora.replace(day=1).date()

    if fecha_hasta_str:
        fecha_hasta = datetime.strptime(fecha_hasta_str, '%Y-%m-%d').date()
    else:
        fecha_hasta = ahora.date()

    query = query.filter(func.date(Llamado.fecha_creacion) >= fecha_desde)
    query = query.filter(func.date(Llamado.fecha_creacion) <= fecha_hasta)
    
    # Ordenamos cronológicamente para el reporte
    llamados = query.order_by(Llamado.fecha_creacion.asc()).all()

    # 2. Carga en bloque de eventos para N+1 queries
    eventos_por_llamado = {}
    if llamados:
        llamados_ids = [ll.id for ll in llamados]
        eventos_bulk = LlamadoEvento.query.filter(
            LlamadoEvento.llamado_id.in_(llamados_ids)
        ).order_by(LlamadoEvento.fecha_evento.asc()).all()

        for ev in eventos_bulk:
            eventos_por_llamado.setdefault(ev.llamado_id, []).append(ev)

    # 3. Crear el libro de Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reporte Operativo"

    # Encabezados
    headers = [
        "ID", "Fecha Creación", "Establecimiento", "Box", "Pacientes", 
        "Estado", "T. Espera (min)", "T. Atención (min)", "Re-llamados", 
        "Creado Por", "Cerrado Por", "Fecha Cierre"
    ]
    ws.append(headers)

    # Estilizar encabezados (Fondo oscuro, letra blanca, centrado)
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    header_align = Alignment(horizontal="center", vertical="center")

    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align

    # MEJORA: Congelar encabezados
    ws.freeze_panes = "A2"

    # 4. Llenar filas
    for row_idx, ll in enumerate(llamados, start=2): # Comenzamos en la fila 2
        box_nombre = ll.box.nombre if ll.box else "N/A"
        est_nombre = ll.establecimiento.nombre if ll.establecimiento else "N/A"
        pacientes_str = ", ".join([p.paciente_nombre for p in ll.pacientes])
        creador_nombre = ll.usuario_creador.nombre_completo if ll.usuario_creador else "N/A"
        cerrador_nombre = ll.usuario_cierre.nombre_completo if ll.usuario_cierre else "N/A"

        # Calcular tiempos específicos para esta fila
        eventos = eventos_por_llamado.get(ll.id, [])
        t_primer_llamado = None
        rellamados = 0

        for ev in eventos:
            if ev.tipo_evento == 'PRIMER_LLAMADO':
                t_primer_llamado = ev.fecha_evento
            elif ev.tipo_evento in ['SEGUNDO_LLAMADO', 'TERCER_LLAMADO']:
                rellamados += 1

        espera_min = ""
        atencion_min = ""

        if t_primer_llamado:
            espera_min = round((t_primer_llamado - ll.fecha_creacion).total_seconds() / 60.0, 1)
            if ll.fecha_cierre and ll.estado == 'FINALIZADO':
                atencion_min = round((ll.fecha_cierre - t_primer_llamado).total_seconds() / 60.0, 1)

        # MEJORA: Insertar Fechas nativas
        row = [
            ll.id,
            ll.fecha_creacion.replace(tzinfo=None) if ll.fecha_creacion else None,
            est_nombre,
            box_nombre,
            pacientes_str,
            ll.estado,
            espera_min,
            atencion_min,
            rellamados,
            creador_nombre,
            cerrador_nombre,
            ll.fecha_cierre.replace(tzinfo=None) if ll.fecha_cierre else None
        ]
        ws.append(row)

        # Formato nativo para fechas en columnas B y L
        if ws.cell(row=row_idx, column=2).value:
            ws.cell(row=row_idx, column=2).number_format = 'YYYY-MM-DD HH:MM:SS'
        if ws.cell(row=row_idx, column=12).value:
            ws.cell(row=row_idx, column=12).number_format = 'YYYY-MM-DD HH:MM:SS'

    # MEJORA: Excel auto filters
    ws.auto_filter.ref = ws.dimensions

    # 5. Autoajustar ancho de columnas para que se vea profesional
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter # Obtenemos la letra (A, B, C...)
        for cell in col:
            try:
                # Considerar formato de fecha en la estimación de ancho
                val_to_check = cell.value.strftime('%Y-%m-%d %H:%M:%S') if isinstance(cell.value, datetime) else str(cell.value)
                if len(val_to_check) > max_length:
                    max_length = len(val_to_check)
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column].width = adjusted_width

    # 6. Preparar el archivo para descarga en memoria (sin guardar en disco)
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    nombre_archivo = f"Reporte_Llamados_{fecha_desde.strftime('%Y%m%d')}_a_{fecha_hasta.strftime('%Y%m%d')}.xlsx"

    return send_file(
        output,
        as_attachment=True,
        download_name=nombre_archivo,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )